#!/usr/bin/env python3
"""
Сборка архива подсказок из рабочего файла Excel-индексатора.

Зачем. У Романа в Excel двадцать тысяч записей, и подсказки там прогреты
годами: знакомый человек находится с первых букв. Наша программа помнит
только набранное в ней самой. Замер 13.09.2026 на незнакомой странице —
55 секунд на запись — сделан с пустой памятью. Этот инструмент переносит
память, а не записи: персон, их места и звания, связки «муж — жена», причт,
частоту употребления. Сами записи о рождениях, браках и смертях не переносятся —
их формы в программе ещё не все есть, и для подсказок они не нужны.

Как пользоваться (у Mike, не у Романа):

    python3 db/tools/build_archive.py "путь/к/Индексатор_МК_….xlsm" архив.sqlite

Получившийся архив.sqlite отправляется Роману, он загружает его в программе
кнопкой «Загрузить архив из файла» на экране «Дело». Слияние делает
db/import_archive.sql — только дополняет, набранное не трогает.

Рабочий файл Романа в репозиторий не попадает никогда. Архив — тоже.

Устройство листов «1», «2», «3» (рождения, браки, смерти) одинаковое:
    13 НП, 14 звание, 15 ИОФ — первая персона (отец / жених / умерший)
    19 НП, 20 звание, 21 ИОФ — вторая (мать / невеста / родитель)
    31, 35, 39, 43 — ИОФ следующих персон, за каждым НП и звание
    47/48, 50/51, 53/54 — ИОФ и звание причта
"""

import re
import sqlite3
import sys
import unicodedata
from pathlib import Path

DB_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(DB_DIR))

# Разбор строк листа. Ключ — имя листа; для каждой персоны: (ИОФ, НП, звание,
# роль). Роль даёт пол там, где он известен наверняка.
SLOTS = {
    "1": [(15, 13, 14, "father"), (21, 19, 20, "mother"),
          (31, 32, 33, "god"), (35, 36, 37, "god")],
    "2": [(15, 13, 14, "groom"), (21, 19, 20, "bride"),
          (31, 32, 33, "witness"), (35, 36, 37, "witness"),
          (39, 40, 41, "witness"), (43, 44, 45, "witness")],
    "3": [(15, 13, 14, "deceased"), (21, 19, 20, "relative")],
}
CLERGY = [(47, 48), (50, 51), (53, 54)]
ROLE_SEX = {"father": "М", "mother": "Ж", "groom": "М", "bride": "Ж"}
# Кто чья жена: в рождениях первая персона — отец, вторая — мать; в браках
# жених и невеста.
SPOUSES = {"1": (0, 1), "2": (0, 1)}


def _utf8_stdout() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass


def norm(value) -> str:
    """Совпадает с norm() в build_seed.py и normalize() в main.rs."""
    if not value:
        return ""
    s = unicodedata.normalize("NFC", str(value)).strip().lower().replace("ё", "е")
    return re.sub(r"\s+", " ", s)


def clean(value) -> str:
    """Текст ячейки. Значения без единой буквы — «0», «-», прочерки — это
    пустота, а не название: в архиве Романа так помечены незаполненные ячейки."""
    if value is None:
        return ""
    s = re.sub(r"\s+", " ", str(value)).strip()
    return s if re.search(r"[^\W\d_]", s) else ""


class Harvest:
    """Что собрано из листов. Ключи словарей — нормализованные, значения
    хранят первое встреченное написание и счётчик."""

    def __init__(self):
        self.persons = {}   # (iof_norm, place_norm, rank_norm) -> [iof, place, rank, sex, uses]
        self.spouses = {}   # (husband_norm, wife_norm) -> [wife_iof, wife_place, wife_rank, uses]
        self.clergy = {}    # (iof_norm, rank_norm) -> [iof, rank, uses]
        self.places = {}    # place_norm -> [place, uses]
        self.usage = {}     # (kind, value_norm) -> [value, count]

    def bump(self, kind, value):
        if not value:
            return
        k = (kind, norm(value))
        if k in self.usage:
            self.usage[k][1] += 1
        else:
            self.usage[k] = [value, 1]


def guess_sex(dictionary: sqlite3.Connection, iof: str, role: str):
    """Пол: сначала по роли, где он известен наверняка, иначе по словарю имён —
    тем же способом, что parse_iof() в приложении."""
    if role in ROLE_SEX:
        return ROLE_SEX[role]
    toks = iof.split()
    if not toks or dictionary is None:
        return None
    row = dictionary.execute(
        "SELECT d.gender FROM name_form f JOIN name_dict d ON d.id = f.name_id "
        "WHERE f.kind IN ('name','variant') AND f.form_norm = ? ORDER BY f.priority LIMIT 1",
        (norm(toks[0]),)).fetchone()
    if row and row[0]:
        return row[0]
    if len(toks) > 1:
        row = dictionary.execute(
            "SELECT f.gender FROM name_form f WHERE f.kind LIKE 'patr%' AND f.form_norm = ? "
            "ORDER BY f.priority LIMIT 1", (norm(toks[1]),)).fetchone()
        if row and row[0]:
            return row[0]
    return None


def harvest_rows(sheet: str, rows, dictionary: sqlite3.Connection, h: Harvest) -> int:
    """Собирает память из строк одного листа. rows — последовательности ячеек,
    индексация с нуля (колонка 1 листа = индекс 0). Возвращает число записей."""
    n = 0
    slots = SLOTS[sheet]
    for r in rows:
        # Признак строки — номер записи в первой колонке. Он число, поэтому
        # смотрим на сырую ячейку, а не через clean(): та считает «без букв»
        # пустотой и выкинула бы все строки разом.
        if not r or r[0] is None or str(r[0]).strip() == "":
            continue
        n += 1
        cell = lambda col: clean(r[col - 1]) if len(r) >= col else ""
        found = []
        for iof_c, place_c, rank_c, role in slots:
            iof, place, rank = cell(iof_c), cell(place_c), cell(rank_c)
            if not iof:
                found.append(None)
                continue
            sex = guess_sex(dictionary, iof, role)
            key = (norm(iof), norm(place), norm(rank))
            if key in h.persons:
                h.persons[key][4] += 1
                if not h.persons[key][3] and sex:
                    h.persons[key][3] = sex
            else:
                h.persons[key] = [iof, place or None, rank or None, sex, 1]
            found.append((iof, place, rank, sex))
            if place:
                pk = norm(place)
                if pk in h.places:
                    h.places[pk][1] += 1
                else:
                    h.places[pk] = [place, 1]
                h.bump("place", place)
            if rank:
                h.bump("rank_f" if sex == "Ж" else "rank_m", rank)
            toks = iof.split()
            if toks:
                h.bump("first_name", toks[0])
            if len(toks) > 1:
                h.bump("patronymic", toks[1])
        if sheet in SPOUSES:
            a, b = SPOUSES[sheet]
            if found[a] and found[b]:
                hus, wife = found[a], found[b]
                key = (norm(hus[0]), norm(wife[0]))
                if key in h.spouses:
                    h.spouses[key][3] += 1
                else:
                    h.spouses[key] = [wife[0], wife[1] or None, wife[2] or None, 1]
        for iof_c, rank_c in CLERGY:
            iof, rank = cell(iof_c), cell(rank_c)
            if not iof:
                continue
            key = (norm(iof), norm(rank))
            if key in h.clergy:
                h.clergy[key][2] += 1
            else:
                h.clergy[key] = [iof, rank or None, 1]
            h.bump("rank_clergy", rank)
    return n


def write_archive(h: Harvest, out: Path, source: str) -> dict:
    """Пишет архив: та же схема, что у программы, заполнены только таблицы
    памяти. Схема берётся из db/schema.sql, чтобы import_archive.sql мог
    обращаться к колонкам по тем же именам."""
    if out.exists():
        out.unlink()
    db = sqlite3.connect(out)
    schema = (DB_DIR / "schema.sql").read_text(encoding="utf-8")
    db.executescript(schema)
    db.execute("INSERT INTO setting (key, value) VALUES ('archive_source', ?)", (source,))
    db.execute("INSERT INTO setting (key, value) VALUES ('archive_built', datetime('now'))")
    db.executemany(
        "INSERT INTO person_index (iof, iof_norm, place, rank, gender, uses, last_used_at) "
        "VALUES (?, ?, ?, ?, ?, ?, datetime('now'))",
        [(v[0], k[0], v[1], v[2], v[3], v[4]) for k, v in h.persons.items()])
    db.executemany(
        "INSERT INTO spouse_index (husband_norm, wife_iof, wife_place, wife_rank, uses, last_used_at) "
        "VALUES (?, ?, ?, ?, ?, datetime('now'))",
        [(k[0], v[0], v[1], v[2], v[3]) for k, v in h.spouses.items()])
    db.executemany(
        "INSERT INTO clergy_index (iof, iof_norm, rank, uses, last_used_at) "
        "VALUES (?, ?, ?, ?, datetime('now'))",
        [(v[0], k[0], v[1], v[2]) for k, v in h.clergy.items()])
    db.executemany(
        "INSERT INTO place (name, name_norm, origin) VALUES (?, ?, 'archive')",
        [(v[0], k) for k, v in h.places.items()])
    db.executemany(
        "INSERT INTO usage_stat (kind, scope, scope_key, value, value_norm, count, last_used_at) "
        "VALUES (?, 'global', '', ?, ?, ?, datetime('now'))",
        [(k[0], v[0], k[1], v[1]) for k, v in h.usage.items()])
    db.commit()
    stats = {
        "персон": len(h.persons), "пар муж—жена": len(h.spouses),
        "причта": len(h.clergy), "населённых пунктов": len(h.places),
        "строк статистики": len(h.usage),
    }
    db.close()
    return stats


def main() -> int:
    _utf8_stdout()
    if len(sys.argv) != 3:
        print(__doc__)
        return 1
    xlsm, out = Path(sys.argv[1]), Path(sys.argv[2])
    if not xlsm.exists():
        print(f"Не найден файл: {xlsm}")
        return 1
    import openpyxl
    import build_seed
    import tempfile

    # Словарь имён для определения пола — из поставки, собранной тут же.
    with tempfile.TemporaryDirectory() as tmp:
        seed = Path(tmp) / "seed.sqlite"
        build_seed.build(seed)
        dictionary = sqlite3.connect(seed)
        wb = openpyxl.load_workbook(xlsm, read_only=True, data_only=True)
        h = Harvest()
        total = 0
        for sheet in ("1", "2", "3"):
            if sheet not in wb.sheetnames:
                continue
            n = harvest_rows(sheet, wb[sheet].iter_rows(min_row=3, values_only=True), dictionary, h)
            print(f"лист «{sheet}»: {n} записей")
            total += n
        dictionary.close()
    stats = write_archive(h, out, xlsm.name)
    print(f"\nАрхив: {out} ({out.stat().st_size // 1024} КБ) из {total} записей")
    for k, v in stats.items():
        print(f"  {k:<22} {v:>6}")
    unknown = sum(1 for v in h.persons.values() if not v[3])
    print(f"  персон без пола        {unknown:>6}  (им подсказка покажется всем)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
